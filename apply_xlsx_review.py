"""Apply XLSX human-review verdicts to clean the -3..0 risk pool (Exp10).

The triage queue (P(安全)>0.5 misclassified positives) was human-labeled in an
XLSX with a verdict column (default `real_label`). This script reports the
verdict distribution + diagnostics and drops non-genuine-risk samples from the
training manifest.

Rule (experiments_log 观察12): queue samples are SUSPECTS. DROP anything a human
marked as invisible / not_sure / low_quality (or any configured drop label);
KEEP genuine visible-risk samples. NEVER flip labels, NEVER re-cut windows — the
event is real (AEB fired); flipping teaches "this precursor = safe", and a
post-trigger re-cut re-introduces the braking leak.

Usage:
    # analyze only (no write) — run this first to confirm column semantics
    python apply_xlsx_review.py --xlsx review.xlsx --train_json pool.json \\
        --out_json cleaned.json --dry_run

    # then apply
    python apply_xlsx_review.py --xlsx review.xlsx --train_json pool.json \\
        --out_json cleaned.json
"""

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("Need openpyxl: pip install openpyxl")

# verdict (real_label) values that mean "remove from training". Everything else
# in the queue is treated as KEEP (genuine hard positive). Case-insensitive.
DEFAULT_DROP_LABELS = {"invisible_risk", "invisible", "not_sure", "notsure",
                       "low_quality", "lowquality", "unknown", "delete", "drop"}


def load_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    raw_header = next(it)
    header = [(str(h).strip() if h is not None else f"_col{i}") for i, h in enumerate(raw_header)]
    records = []
    for row in it:
        if row is None or all(c is None for c in row):
            continue
        records.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
    return header, records


def norm(v):
    return str(v).strip().lower() if v is not None else ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--train_json", required=True, help="Risk-pool manifest to clean (e.g. 32k pool)")
    ap.add_argument("--out_json", required=True)
    ap.add_argument("--stem_col", default="stem")
    ap.add_argument("--verdict_col", default="real_label")
    ap.add_argument("--p_safe_col", default="p_safe")
    ap.add_argument("--scene_col", default="", help="Optional condition column for analysis only (empty = skip)")
    ap.add_argument("--annotator_col", default="num_type", help="Optional annotator column for analysis")
    ap.add_argument("--drop_labels", default=None,
                    help="Comma-separated verdict values to DROP (default: the known bad set)")
    ap.add_argument("--dry_run", action="store_true", help="Analyze only, do not write out_json")
    args = ap.parse_args()

    header, recs = load_xlsx(args.xlsx)
    print("=" * 72)
    print(f"📄 {args.xlsx}")
    print(f"   columns: {header}")
    print(f"   rows: {len(recs)}")
    if recs:
        print(f"   sample row: { {k: recs[0].get(k) for k in header[:8]} }")

    for col in (args.stem_col, args.verdict_col):
        if col not in header:
            raise SystemExit(f"❌ column '{col}' not in header {header}; pass the right --{'stem' if col==args.stem_col else 'verdict'}_col")

    drop_labels = (set(s.strip().lower() for s in args.drop_labels.split(",")) if args.drop_labels
                   else set(DEFAULT_DROP_LABELS))

    # ---- analysis: verdict distribution ----
    verdict_counts = Counter(norm(r.get(args.verdict_col)) for r in recs)
    print("\n" + "=" * 72)
    print(f"🏷  verdict 分布（列 '{args.verdict_col}'）—— DROP 标记按 --drop_labels")
    print("=" * 72)
    n = len(recs)
    n_empty = verdict_counts.get("", 0)
    for v, c in verdict_counts.most_common():
        tag = "（空，未标注）" if v == "" else ("→ DROP" if v in drop_labels else "→ KEEP")
        print(f"  {v or '<空>':<22} {c:>6}  {c/max(n,1):>6.1%}  {tag}")
    keep_n = sum(c for v, c in verdict_counts.items() if v and v not in drop_labels)
    drop_n = sum(c for v, c in verdict_counts.items() if v in drop_labels)
    print(f"\n  KEEP(可见风险，留): {keep_n}   DROP(剔除): {drop_n}   空/未标注: {n_empty}")
    if n_empty:
        print(f"  ⚠️ {n_empty} 行 verdict 为空（未标注）→ 默认按 KEEP 处理（不剔除），请确认是否漏标")
    print(f"  → 队列 keep 率 = {keep_n/max(n,1):.1%}（模型自信判错却确为真风险的比例；1-此 = 噪声率）")

    # ---- analysis: p_safe by verdict (validates the probe ranking) ----
    if args.p_safe_col in header:
        by_v = defaultdict(list)
        for r in recs:
            try:
                by_v[norm(r.get(args.verdict_col))].append(float(r[args.p_safe_col]))
            except (TypeError, ValueError, KeyError):
                pass
        print(f"\n  --- P(安全) by verdict（探针有效性：invisible 应更高）---")
        for v in sorted(by_v, key=lambda k: -sum(by_v[k])/len(by_v[k])):
            xs = by_v[v]
            print(f"  {v or '<空>':<22} n={len(xs):>5}  mean={sum(xs)/len(xs):.3f}  "
                  f"min={min(xs):.3f}  max={max(xs):.3f}")

    # ---- analysis: scene/condition cross-tab (only if explicitly requested) ----
    if args.scene_col and args.scene_col in header:
        print(f"\n  --- '{args.scene_col}' × verdict（哪类条件多 DROP）---")
        cross = defaultdict(Counter)
        for r in recs:
            cross[norm(r.get(args.scene_col)) or "<空>"][norm(r.get(args.verdict_col))] += 1
        for scene, vc in sorted(cross.items(), key=lambda kv: -sum(kv[1].values()))[:15]:
            d = sum(c for v, c in vc.items() if v in drop_labels)
            tot = sum(vc.values())
            print(f"  {scene:<22} 总{tot:>4}  DROP {d:>4} ({d/max(tot,1):>5.0%})")

    # ---- analysis: annotator load ----
    if args.annotator_col in header:
        ac = Counter(norm(r.get(args.annotator_col)) or "<空>" for r in recs)
        print(f"\n  --- 标注者负载（列 '{args.annotator_col}'）---")
        for a, c in ac.most_common():
            print(f"  {a:<16} {c}")

    # ---- build drop set ----
    drop_stems = set()
    for r in recs:
        if norm(r.get(args.verdict_col)) in drop_labels:
            stem = str(r.get(args.stem_col)).strip()
            if stem and stem.lower() != "none":
                drop_stems.add(Path(stem).stem)  # tolerate path or bare stem

    # ---- apply to manifest ----
    with open(args.train_json, "r", encoding="utf-8") as f:
        train = json.load(f)

    def label_of(s):
        a = next(m for m in s["messages"] if m.get("role") == "assistant")
        raw = str(a["content"]).strip()
        return 1 if raw in ("1", "高风险") else (0 if raw in ("0", "安全") else int(raw))

    seen = set()
    kept, dropped = [], 0
    for s in train:
        vids = s.get("videos") or []
        st = Path(vids[0]).stem if vids else None
        if st in drop_stems:
            dropped += 1
            seen.add(st)
        else:
            kept.append(s)
    not_found = drop_stems - seen

    pos = sum(1 for s in kept if label_of(s) == 1)
    neg = len(kept) - pos
    print("\n" + "=" * 72)
    print("📦 清洗结果")
    print("=" * 72)
    print(f"  manifest 原始 {len(train)}  → 剔除 {dropped}  → 剩余 {len(kept)}")
    print(f"  剩余 正(高风险) {pos} : 负(安全) {neg} = {pos/max(neg,1):.2f}:1")
    print(f"  DROP 判定共 {len(drop_stems)} 个 stem，其中 {len(not_found)} 个不在本 manifest "
          f"（预期=备用池/非本集样本）")

    if args.dry_run:
        print("\n  [dry_run] 未写出。确认列与映射无误后去掉 --dry_run 再跑。")
    else:
        with open(args.out_json, "w", encoding="utf-8") as f:
            json.dump(kept, f, ensure_ascii=False, indent=2)
        print(f"\n  ↳ {args.out_json}（仅剔除，无翻标签/换窗口）")


if __name__ == "__main__":
    main()
